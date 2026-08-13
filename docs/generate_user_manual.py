"""Generate the AI Blasting Suite Microsoft Word user manual."""

from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "AI_Blasting_Suite_User_Manual.docx"
DEPLOYMENT_URL = "https://ai-blasting-kvzex3f0i-navan-labs.vercel.app"

BLUE = "1D4ED8"
DARK = "172033"
ORANGE = "EA580C"
PALE_BLUE = "EFF6FF"
PALE_ORANGE = "FFF7ED"
PALE_YELLOW = "FEF3C7"
PALE_RED = "FEF2F2"
GREY = "64748B"
LIGHT_GREY = "E2E8F0"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_field(paragraph, field: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def add_toc(document: Document) -> None:
    sections = [
        "1. About this manual",
        "2. Access, sign-in and sign-out",
        "3. Workspace orientation",
        "4. Data Manager",
        "5. Prediction",
        "6. Feature Importance & Explainable AI",
        "7. Parameter Optimisation",
        "8. Cost Optimisation",
        "9. GeoMotion 3D",
        "10. Slope Stability",
        "11. Back Break",
        "12. Flyrock (ML + Empirical)",
        "13. File formats and output handling",
        "14. Troubleshooting",
        "15. Glossary",
        "16. Owner completion checklist",
    ]
    table = document.add_table(rows=8, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, entry in enumerate(sections):
        row = index % 8
        column = index // 8
        cell = table.cell(row, column)
        set_cell_margins(cell, top=65, start=100, bottom=65, end=100)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(entry)
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string(BLUE)


def add_table(document: Document, headers: list[str], rows: Iterable[Iterable[object]], widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Manual Table"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    header = table.rows[0]
    set_repeat_table_header(header)
    for idx, text in enumerate(headers):
        cell = header.cells[idx]
        set_cell_shading(cell, BLUE)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(str(text))
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(WHITE)
        run.font.size = Pt(9)
        if widths:
            cell.width = widths[idx]
    for row_data in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_data):
            cell = cells[idx]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if len(table.rows) % 2 == 1:
                set_cell_shading(cell, "F8FAFC")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(str(value))
            run.font.size = Pt(9)
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_callout(document: Document, title: str, text: str, kind: str = "note") -> None:
    fills = {"note": PALE_BLUE, "warning": PALE_ORANGE, "action": PALE_YELLOW, "danger": PALE_RED}
    colors = {"note": BLUE, "warning": ORANGE, "action": "92400E", "danger": "B91C1C"}
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fills[kind])
    set_cell_margins(cell, top=140, start=180, bottom=140, end=180)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(title)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(colors[kind])
    run.font.size = Pt(10)
    paragraph = cell.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.style = document.styles["Body Text"]
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_bullets(document: Document, items: Iterable[str], level: int = 0) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        paragraph.add_run(item)


def add_steps(document: Document, items: Iterable[str]) -> None:
    for index, item in enumerate(items, start=1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.5)
        paragraph.paragraph_format.first_line_indent = Cm(-0.5)
        number = paragraph.add_run(f"{index}. ")
        number.bold = True
        paragraph.add_run(item)


def add_screenshot(
    document: Document,
    caption: str,
    keywords: tuple[str, ...],
    owner_instruction: str,
    width: float = 6.35,
) -> None:
    del keywords, width
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F8FAFC")
    set_cell_margins(cell, top=420, start=180, bottom=420, end=180)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label = paragraph.add_run("FIGURE PLACEHOLDER")
    label.bold = True
    label.font.color.rgb = RGBColor.from_string(BLUE)
    label.font.size = Pt(11)
    detail = cell.add_paragraph(owner_instruction)
    detail.alignment = WD_ALIGN_PARAGRAPH.CENTER
    detail.runs[0].font.color.rgb = RGBColor.from_string(GREY)
    caption_paragraph = document.add_paragraph(caption, style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_module_intro(document: Document, purpose: str, use_when: str, outputs: str) -> None:
    add_table(
        document,
        ["Purpose", "Use this module when", "Primary outputs"],
        [[purpose, use_when, outputs]],
    )


def configure_document() -> Document:
    document = Document()
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    document.settings._element.append(update_fields)
    section = document.sections[0]
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    body = styles["Body Text"]
    body.font.name = "Aptos"
    body.font.size = Pt(9.5)
    body.font.color.rgb = RGBColor.from_string(DARK)
    body.paragraph_format.space_after = Pt(5)
    body.paragraph_format.line_spacing = 1.08

    for name, size, color in (
        ("Title", 34, DARK),
        ("Heading 1", 22, DARK),
        ("Heading 2", 15, BLUE),
        ("Heading 3", 11, ORANGE),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    styles["Title"].paragraph_format.space_before = Pt(0)
    styles["Title"].paragraph_format.space_after = Pt(10)

    caption = styles["Caption"]
    caption.font.name = "Aptos"
    caption.font.size = Pt(8.5)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(GREY)

    if "Manual Table" not in styles:
        table_style = styles.add_style("Manual Table", WD_STYLE_TYPE.TABLE)
        table_style.base_style = styles["Table Grid"]

    header = section.header.paragraphs[0]
    header.text = "AI BLASTING SUITE  /  USER MANUAL"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.runs[0]
    header_run.font.name = "Aptos"
    header_run.font.size = Pt(8)
    header_run.font.bold = True
    header_run.font.color.rgb = RGBColor.from_string(GREY)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Controlled copy when viewed electronically  •  Page ")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(GREY)
    add_page_field(footer, "PAGE")
    footer.add_run(" of ")
    add_page_field(footer, "NUMPAGES")
    return document


def build_manual() -> Document:
    document = configure_document()

    # Cover
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(72)
    run = paragraph.add_run("AI")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor.from_string(ORANGE)
    run = paragraph.add_run(" BLASTING SUITE")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    title = document.add_paragraph("User Manual", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle = document.add_paragraph(
        "Operating instructions, module purposes, inputs, outputs, formats and units"
    )
    subtitle.style = document.styles["Subtitle"]
    subtitle.runs[0].font.color.rgb = RGBColor.from_string(GREY)

    document.add_paragraph("\n")
    add_table(
        document,
        ["Document field", "Value"],
        [
            ["Software", "AI Blasting Suite — React/FastAPI web application"],
            ["Deployment", DEPLOYMENT_URL],
            ["Manual version", "1.1 — GeoMotion 3D deployment"],
            ["Prepared", date.today().strftime("%d %B %Y")],
            ["Intended users", "Blasting engineers, reviewers, planners, analysts, trainers and authorised students"],
            ["Status", "Draft for owner review and site-specific completion"],
        ],
    )
    add_callout(
        document,
        "SAFETY-CRITICAL SCOPE",
        "This software supports analysis, planning, education and simulation. It does not arm, fire, program or communicate with detonator hardware. Every design, prediction and export must be checked and approved by qualified blasting personnel under applicable legislation, mine procedures, manufacturer requirements and a site-specific risk assessment.",
        "danger",
    )
    add_callout(
        document,
        "PROPRIETARY USE NOTICE",
        "This application is proprietary Debswana software. Use and distribution are restricted to authorised internal business purposes and remain subject to the repository licence and organisational information-governance requirements.",
        "warning",
    )
    add_callout(
        document,
        "OWNER REVIEW REQUIRED",
        "Yellow boxes identify owner-supplied governance information that still requires completion. Labelled figure placeholders are intentionally empty so the document owner can insert approved deployment screenshots.",
        "action",
    )
    document.add_page_break()

    document.add_heading("Document control and approval", level=1)
    add_table(
        document,
        ["Role", "Name", "Signature", "Date"],
        [
            ["Document owner", "[OWNER TO COMPLETE]", "", ""],
            ["Technical reviewer", "[OWNER TO COMPLETE]", "", ""],
            ["Qualified blasting reviewer", "[OWNER TO COMPLETE]", "", ""],
            ["Release approver", "[OWNER TO COMPLETE]", "", ""],
        ],
    )
    document.add_heading("Revision history", level=2)
    add_table(
        document,
        ["Version", "Date", "Author/owner", "Change"],
        [
            ["1.0", date.today().isoformat(), "Generated from repository baseline", "Initial comprehensive user manual"],
            ["1.1", date.today().isoformat(), "Updated from deployed GeoMotion branch", "Set production deployment, add GeoMotion 3D and replace screenshots with figure labels"],
        ],
    )
    add_callout(
        document,
        "OWNER ACTION REQUIRED — DOCUMENT METADATA",
        "Add the organisation name/logo, document identifier, software release/build number, support contact, formal approval names and next review date. Verify that the deployment URL remains current.",
        "action",
    )

    document.add_heading("Contents", level=1)
    add_toc(document)
    document.add_page_break()

    # 1
    document.add_heading("1. About this manual", level=1)
    document.add_heading("1.1 Purpose", level=2)
    document.add_paragraph(
        f"This manual explains how authorised users sign in, prepare data, run each analysis or planning module, interpret its outputs and export results. It covers the deployment at {DEPLOYMENT_URL}, based on the GeoMotion 3D deployment branch."
    )
    document.add_heading("1.2 What the suite does", level=2)
    document.add_paragraph(
        "AI Blasting Suite combines dataset quality checks, empirical calculations, machine-learning predictions, explainability, design optimisation, safety screening and three-dimensional blast-movement analysis in one workspace. Its modules share the active combined dataset where indicated."
    )
    document.add_heading("1.3 Intended audience", level=2)
    add_bullets(
        document,
        [
            "Blasting engineers and licensed blasting personnel performing qualified review.",
            "Technical services, planning and operations personnel comparing design scenarios.",
            "Geotechnical and safety personnel screening slope, flyrock and back-break risk.",
            "Data analysts maintaining datasets and reviewing model behaviour.",
            "Students and trainers using the planning and simulation functions under supervision.",
        ],
    )
    document.add_heading("1.4 Safety and decision limitations", level=2)
    add_bullets(
        document,
        [
            "Predictions and optimised recipes are decision-support outputs, not approvals or firing instructions.",
            "Machine-learning quality depends on the relevance, size and quality of the training dataset.",
            "Empirical constants and acceptance limits must be calibrated for the site.",
            "GeoMotion 3D is an uncalibrated planning and study tool. It cannot control explosives or initiation hardware and must not be used for field execution, dig-limit control, resource reporting or production decisions.",
            "A low predicted risk or a model classification never replaces engineering judgement, inspections, statutory controls or exclusion-zone procedures.",
        ],
    )
    document.add_heading("1.5 Conventions and common units", level=2)
    document.add_paragraph(
        "Figures are intentionally labelled placeholders. The document owner can insert approved screenshots from the deployment without changing the operating instructions. Values shown in examples are demonstration values unless their provenance is explicitly identified as measured."
    )
    add_table(
        document,
        ["Symbol/term", "Meaning", "Unit or format"],
        [
            ["B / Burden", "Distance from a blasthole row to the free face or preceding row", "m"],
            ["S / Spacing", "Distance between holes in a row", "m"],
            ["HPD", "Holes firing in the same delay group", "count (dimensionless)"],
            ["PPV / Ground vibration", "Peak particle velocity", "mm/s"],
            ["Airblast", "Air overpressure level", "dB"],
            ["X50", "Fragment size at 50% passing", "mm"],
            ["Xm", "Kuz–Ram mean fragmentation size", "mm"],
            ["PF", "Powder factor", "kg/m³"],
            ["RWS", "Relative weight strength", "dimensionless relative value"],
            ["R²", "Model coefficient of determination", "dimensionless; higher is generally better"],
            ["Probability", "Model likelihood or screening score", "0–1 internally; often displayed as %"],
            ["Delay", "Time from initiation reference", "ms"],
            ["Coordinates", "X/easting, Y/northing and optional Z/elevation", "Project coordinate units; normally m, confirm site system"],
            ["Cost", "Initiation, explosive, drilling and total cost", "BWP in the current interface"],
            ["cpht", "Diamond grade", "carats per hundred tonnes"],
            ["VOD", "Velocity of detonation", "m/s"],
            ["Voxel", "Three-dimensional material cell", "1 m × 1 m × 1 m in the authoritative GeoMotion model"],
        ],
    )

    # 2
    document.add_heading("2. Access, sign-in and sign-out", level=1)
    document.add_heading("2.1 Current configured access list", level=2)
    document.add_paragraph(
        "The following addresses are the default allowlist found in both the frontend and backend configuration:"
    )
    add_table(
        document,
        ["Authorised email"],
        [
            ["so13000604@biust.ac.bw"],
            ["Ozigwa@debswana.bw"],
            ["Tgalefete@debswana.bw"],
            ["Mhiya@debswana.bw"],
            ["Ttshambane@debswana.bw"],
            ["Mgaopelo@debswana.bw"],
            ["SMoabi@debswana.bw"],
            ["MMoleofe@debswana.bw"],
        ],
    )
    add_callout(
        document,
        "CONFIGURATION NOTE",
        "Deployment configuration can add frontend addresses or replace backend defaults through VITE_ALLOWED_LOGIN_EMAILS and BLAST_ALLOWED_AUTH_EMAILS. Treat the table above as the deployment-branch baseline, not proof of the live production allowlist.",
        "warning",
    )
    document.add_paragraph(
        "The current application implements an email allowlist rather than separate user roles. An authorised user can navigate to the same user-facing modules; organisational approvals and operating authority remain outside the software."
    )
    add_callout(
        document,
        "OWNER ACTION REQUIRED — VERIFY ACCESS",
        "Before issue, compare this table with the current production environment allowlist and the organisation's access approval record. Add an access owner and review frequency. Do not add passwords or verification codes.",
        "action",
    )
    document.add_heading("2.2 Sign in with a one-time code", level=2)
    add_steps(
        document,
        [
            f"Open {DEPLOYMENT_URL} in a current browser. Vercel Authentication may require an authorised Vercel session before the application sign-in page is displayed.",
            "Enter an authorised work email address.",
            "Select Send magic code.",
            "Retrieve the one-time code from the mailbox and enter it in Verification code.",
            "Select Verify and continue. If the email or code is rejected, see Troubleshooting.",
        ],
    )
    add_screenshot(
        document,
        "Figure 1 — Passwordless sign-in screen.",
        ("login",),
        "Insert a production sign-in screenshot that contains no real one-time code, mailbox content or browser-stored token.",
    )
    document.add_heading("2.3 Session and sign-out", level=2)
    document.add_paragraph(
        "The browser stores an InstantDB refresh token and the signed-in email in local storage so a session can be restored after refresh. Use Sign out in the application header when finished, especially on a shared workstation. Do not copy browser storage values or authentication tokens into support requests."
    )

    # 3
    document.add_heading("3. Workspace orientation", level=1)
    document.add_heading("3.1 Navigation groups", level=2)
    add_table(
        document,
        ["Navigation group", "Modules"],
        [
            ["Analysis", "Prediction; Feature Importance; Parameter Optimisation"],
            ["Operations", "Cost Optimisation; GeoMotion 3D"],
            ["Safety / Geo", "Slope Stability; Back Break; Flyrock"],
            ["Admin", "Data Manager"],
        ],
    )
    document.add_heading("3.2 Header controls", level=2)
    add_bullets(
        document,
        [
            "Dataset selects the shared combined dataset used by Data Manager defaults, Prediction, Feature Importance and Parameter Optimisation.",
            "Theme selects System, Light or Dark appearance.",
            "Accent selects blue, green or dark-blue interface accents.",
            "Backend status indicates whether the service responds. Analysis requiring the backend will fail when it is unavailable.",
            "The user email and Sign out control identify and end the current session.",
        ],
    )
    document.add_heading("3.3 Welcome dashboard", level=2)
    add_bullets(
        document,
        [
            "Current user, last app use and last module used summarise recent activity.",
            "Module guide cards explain each module and provide Open module shortcuts.",
            "Quick access opens Prediction, Cost Optimisation or Flyrock.",
            "Export professional report creates printable HTML containing the active dataset, module guide, generating user and recent user/module/timestamp activity. Review this personal and operational metadata before distributing the report.",
            "Recent activity is a convenience summary, not a complete security or statutory audit log.",
        ],
    )
    add_callout(
        document,
        "DATASET CHANGE EFFECT",
        "Changing the combined dataset resets prediction and optimisation model caches. The next request may take longer while models are prepared again.",
        "note",
    )
    add_screenshot(
        document,
        "Figure 2 — Welcome dashboard and main navigation.",
        ("home", "welcome", "dashboard"),
        "Insert the Welcome dashboard with the sidebar, dataset selector, backend status and module guide visible.",
    )
    document.add_heading("3.4 Recommended end-to-end workflow", level=2)
    add_steps(
        document,
        [
            "Confirm that the correct combined dataset is selected.",
            "Use Data Manager to inspect ranges, missing values, correlations and compliance filters.",
            "Run Prediction for the proposed design and compare empirical and ML results.",
            "Use Feature Importance to understand influential variables and data structure.",
            "Use Parameter Optimisation or Cost Optimisation to compare scenarios and trade-offs.",
            "Use Slope Stability, Back Break and Flyrock for additional safety screening.",
            "Use GeoMotion 3D to validate a delay-bearing charged-hole tie-up and study synthetic post-blast material movement, recovery, loss and dilution.",
            "Export results, record dataset/version details and obtain qualified review.",
        ],
    )

    # 4 Data Manager
    document.add_heading("4. Data Manager", level=1)
    add_module_intro(
        document,
        "Load, inspect, append, filter, visualise, calibrate and export tabular site datasets.",
        "Preparing or quality-checking data before prediction, explainability or optimisation.",
        "Filtered CSV, nominal Excel export, site-model JSON, statistics, audits and plots.",
    )
    document.add_heading("4.1 Input format", level=2)
    add_table(
        document,
        ["Requirement", "Details"],
        [
            ["Accepted file", "Comma-separated values (.csv)"],
            ["Header", "First non-empty row; unique descriptive column names recommended"],
            ["Values", "One observation per row; numeric analysis columns should contain decimal numbers without unit text"],
            ["Append behaviour", "Creates the union of both files' columns and appends rows; missing fields remain blank"],
            ["Shared effect", "The loaded data becomes the active in-browser dataset for Prediction, Feature Importance and Parameter Optimisation"],
        ],
    )
    document.add_heading("4.2 Controls and outputs", level=2)
    add_table(
        document,
        ["Tab/control", "How to use it", "Output"],
        [
            ["Load CSV", "Replace the current in-browser data with a CSV.", "Table and row/column count"],
            ["Append CSV", "Add rows from another CSV; review column alignment.", "Merged table"],
            ["Table", "Search text across all fields and add rows.", "Filtered display or added row"],
            ["Summary", "Review descriptive statistics and enter PPV limit, air limit and HPD.", "Statistics and quick compliance audits"],
            ["Visuals", "Choose plot type and X/Y columns; optional logarithmic axes.", "Scatter, line, bar, histogram, box, hexbin or engineering plot"],
            ["Correlations", "Review numeric-column relationships.", "Correlation heatmap"],
            ["Filters", "Enter numeric minimum/maximum values or a row query.", "Filtered rows"],
            ["Calibration", "Enter HPD and run PPV, airblast or fragmentation calibration.", "Calibration log and site_model.json"],
            ["Export", "Export the currently filtered rows.", "filtered.csv or nominal filtered.xlsx"],
        ],
    )
    add_callout(
        document,
        "CURRENT FORMAT LIMITATION",
        "The control labelled “Export Filtered → Excel” currently writes comma-separated text but names the file filtered.xlsx. Use the CSV export for a reliable file, or open the exported content as CSV and save it as a true Excel workbook. The software owner should correct or validate this behaviour before relying on the Excel label.",
        "warning",
    )
    document.add_heading("4.3 Procedure", level=2)
    add_steps(
        document,
        [
            "Select Load CSV and choose the approved dataset.",
            "Confirm the displayed row and column counts.",
            "Review Table and Summary for blanks, implausible values and unit consistency.",
            "Use Visuals and Correlations to identify outliers and strongly related columns.",
            "Apply Filters or a query to create the required subset.",
            "If calibration is required, enter HPD and run only the calibration supported by suitable measured outputs.",
            "Export the filtered CSV and record the source file, filter criteria and date.",
        ],
    )
    add_callout(
        document,
        "QUERY EXAMPLE",
        "Use backticks around column names containing spaces, for example: `Ground Vibration` <= 12.5 and `Airblast` <= 134. Queries use comparison expressions; use only approved expressions and never paste untrusted code into this field.",
        "note",
    )
    add_screenshot(
        document,
        "Figure 3 — Data Manager table and tabs.",
        ("data", "manager", "table", "loaded"),
        "Insert Data Manager after loading an approved non-sensitive sample CSV, showing row/column count and the main tabs.",
    )
    add_screenshot(
        document,
        "Figure 3A — Data Manager descriptive statistics and quick audits.",
        ("data", "manager", "summary"),
        "Insert the Data Manager Summary tab with approved example limits and no sensitive site data.",
    )
    add_screenshot(
        document,
        "Figure 3B — Data Manager correlation heatmap.",
        ("data", "manager", "correlations"),
        "Insert the Data Manager Correlations tab using an approved sample dataset.",
    )

    # 5 Prediction
    document.add_heading("5. Prediction", level=1)
    add_module_intro(
        document,
        "Estimate ground vibration, airblast and fragmentation for one blast design using empirical equations and optional machine learning.",
        "Comparing a proposed design with site limits and checking empirical/ML agreement.",
        "Output comparison chart, Rosin–Rammler curve, alerts, prediction_result.csv and printable HTML report.",
    )
    document.add_heading("5.1 Blast inputs", level=2)
    add_table(
        document,
        ["Input", "Purpose", "Unit"],
        [
            ["Hole depth", "Drilled hole length used for charge and geometry calculations.", "m"],
            ["Hole diameter", "Nominal blasthole diameter.", "mm"],
            ["Burden", "Distance to free face/preceding row.", "m"],
            ["Spacing", "Hole spacing within a row.", "m"],
            ["Stemming", "Uncharged collar length.", "m"],
            ["Distance", "Distance from blast to monitoring/receptor point.", "m"],
            ["Powder factor", "Explosive mass per blast volume.", "kg/m³"],
            ["Rock density", "Bulk rock density.", "t/m³"],
            ["Linear charge", "Explosive mass per charged metre.", "kg/m"],
            ["Explosive mass", "Explosive mass represented by the row/design.", "kg"],
            ["Blast volume", "Rock volume represented by the design.", "m³"],
            ["# Holes", "Number of blastholes.", "count"],
        ],
    )
    document.add_heading("5.2 Empirical and RR settings", level=2)
    add_table(
        document,
        ["Input", "Default", "Purpose"],
        [
            ["K_ppv", "1000", "Site vibration constant in the scaled-distance PPV equation"],
            ["β", "1.60", "Vibration decay exponent"],
            ["K_air", "170", "Airblast intercept/calibration constant"],
            ["B_air", "20", "Airblast logarithmic slope"],
            ["A_kuz", "22", "Kuz–Ram rock factor"],
            ["RWS", "115", "Relative explosive strength"],
            ["HPD", "1", "Holes contributing to charge per delay"],
            ["Oversize threshold", "500 mm", "Size above which RR oversize percentage is calculated"],
            ["RR n mode", "Estimate", "Estimate uniformity n or use a manual value"],
            ["Manual n", "1.8", "Rosin–Rammler uniformity parameter when Manual is selected"],
            ["Use ML", "Yes", "Include available ML predictions; No uses empirical outputs only"],
        ],
    )
    document.add_heading("5.3 Outputs and interpretation", level=2)
    add_table(
        document,
        ["Output", "Unit", "Interpretation"],
        [
            ["Ground Vibration / PPV", "mm/s", "Compare with the approved site/receptor limit."],
            ["Airblast", "dB", "Compare with the approved site/receptor limit."],
            ["Fragmentation", "mm", "Predicted characteristic size; confirm which statistic the selected dataset uses."],
            ["RR n", "dimensionless", "Higher values indicate a narrower particle-size distribution."],
            ["Xm / X20 / X50 / X80", "mm", "Characteristic fragmentation sizes from the RR curve."],
            ["Oversize", "%", "Predicted fraction above the selected oversize threshold."],
            ["Threshold log", "text", "Flags outputs that exceed user-entered thresholds; checks ML first when available."],
        ],
    )
    document.add_heading("5.4 Procedure", level=2)
    add_steps(
        document,
        [
            "Confirm the dataset source shown at the top of the module.",
            "Enter each blast input, or select Use Medians to initialise values from the current data.",
            "Enter site-approved empirical constants, HPD and RR settings.",
            "Enter threshold values for any required alerts.",
            "Select Predict (ML + Empirical).",
            "Compare ML and empirical outputs; investigate material disagreement instead of choosing the preferred answer.",
            "Open Fragmentation (RR Curve) and review X20/X50/X80 and oversize.",
            "Export the CSV or professional report and retain the dataset and constants with the result.",
        ],
    )
    add_screenshot(
        document,
        "Figure 4 — Prediction inputs and empirical/ML output comparison.",
        ("prediction", "results", "chart"),
        "Insert Prediction after a successful sample run. Show the dataset source, key inputs and empirical-versus-ML output chart without sensitive operational data.",
    )
    add_screenshot(
        document,
        "Figure 4A — Prediction Rosin–Rammler fragmentation curve.",
        ("prediction", "fragmentation", "curve"),
        "Insert the Prediction Fragmentation (RR Curve) tab after a successful approved sample run.",
    )

    # 6 Feature Importance
    document.add_heading("6. Feature Importance & Explainable AI", level=1)
    add_module_intro(
        document,
        "Explain which inputs drive model outputs and reveal correlation or low-dimensional structure in the active dataset.",
        "Prioritising design variables, checking redundant inputs or reviewing model behaviour after prediction.",
        "RF importance, permutation sensitivity, local impact waterfall, partial dependence, correlation map and PCA plots.",
    )
    document.add_heading("6.1 Inputs", level=2)
    add_bullets(
        document,
        [
            "Active shared or uploaded dataset from Data Manager.",
            "Top-K features: integer from 5 to 30; default 12.",
            "Explainability output: select the model output to inspect after computation.",
        ],
    )
    document.add_heading("6.2 Procedure and interpretation", level=2)
    add_steps(
        document,
        [
            "Confirm the dataset name and source shown in the module.",
            "Choose Top-K and select Compute RF Importance.",
            "Check mapping mode, rows used and rows dropped. Positional split means the application could not map standard names and assumed the final three columns are outputs.",
            "Compare global RF importance with permutation sensitivity. A variable that is strong in both views is generally more credible than one strong in only one metric.",
            "Select an output and review local impacts and partial dependence for direction, not causation.",
            "Run PCA Analysis to review explained variance, PC1/PC2 separation and loading variables.",
            "Use the findings to choose optimisation axes, but retain engineering constraints and site knowledge.",
        ],
    )
    add_callout(
        document,
        "INTERPRETATION CAUTION",
        "Importance and correlation do not prove causation. Correlated features can share or mask importance, and extrapolation outside the dataset range is unreliable.",
        "warning",
    )
    add_screenshot(
        document,
        "Figure 5 — Feature importance and explainability results.",
        ("feature", "importance", "fragmentation"),
        "Insert a successful Feature Importance run showing mapping diagnostics and at least one ranking or explainability chart.",
    )

    # 7 parameter
    document.add_heading("7. Parameter Optimisation", level=1)
    add_module_intro(
        document,
        "Build a physics-informed surrogate from the active dataset, search Pareto candidate recipes and perform inverse design.",
        "Exploring design trade-offs or finding inputs likely to achieve a target output.",
        "Pareto candidates, 2D/3D surfaces, selected recipes, goal-seek recipe, CSV surface and printable HTML report.",
    )
    document.add_heading("7.1 Controls", level=2)
    add_table(
        document,
        ["Control", "Meaning"],
        [
            ["Output", "Dataset output to maximise, minimise or target; its unit is inherited from that column."],
            ["Input 1 / Input 2", "Variables shown on surface X and Y axes; their units are inherited from the dataset headers."],
            ["Maximise / Minimise", "Direction applied to the selected output."],
            ["Run Pareto Optimisation", "Runs one search and stores candidate recipes."],
            ["Target output", "Desired value for inverse design, in the selected output's unit."],
            ["Tolerance", "Maximum acceptable absolute target error, in the same output unit."],
            ["Run Goal Seek", "Searches all inputs for a recipe near the target."],
        ],
    )
    document.add_heading("7.2 Procedure", level=2)
    add_steps(
        document,
        [
            "Confirm the active dataset and verify input/output units from its headers.",
            "Select the output, two different input axes and Maximise or Minimise.",
            "Select Run Pareto Optimisation and wait for candidate recipes.",
            "Review model train/test R², fragmentation target band and candidate count.",
            "Click or hover the surface to inspect the complete recipe and all predicted outputs.",
            "Change the display axes to reproject the saved run; this does not rerun the optimiser.",
            "For inverse design, enter target and tolerance and select Run Goal Seek (All Inputs).",
            "Export the surface CSV or optimisation report and obtain engineering review.",
        ],
    )
    add_callout(
        document,
        "UNIT RULE",
        "This module does not impose one fixed unit dictionary. It inherits names and units from the active dataset. Include units in dataset headers and never combine columns with inconsistent units.",
        "note",
    )
    add_screenshot(
        document,
        "Figure 6 — Parameter Optimisation module and run controls.",
        ("parameter", "optimisation"),
        "Insert a completed Parameter Optimisation run showing objective, axes, optimisation summary and the surface explorer.",
    )

    # 8 cost
    document.add_heading("8. Cost Optimisation", level=1)
    add_module_intro(
        document,
        "Compute blast cost and engineering KPIs, optimise geometry under constraints and explore non-dominated cost/quality/safety trade-offs.",
        "Balancing budget, fragmentation, vibration and airblast for candidate designs.",
        "Engineering report, KPI cards, cost/penalty charts, constraint checks, Pareto CSVs and printable HTML report.",
    )
    document.add_heading("8.1 Input groups and defaults", level=2)
    add_table(
        document,
        ["Input", "Default", "Unit / purpose"],
        [
            ["Diameter", "102", "mm"],
            ["Bench height / Burden / Spacing / Subdrilling / Stemming", "10 / 3 / 3.3 / 2 / 1.8", "m"],
            ["Number of holes / HPD", "30 / 1", "counts"],
            ["Block volume", "0", "m³; 0 calculates B × S × bench × holes"],
            ["Explosive density", "1.15", "g/cc"],
            ["RWS", "115", "relative strength"],
            ["Initiation / explosive / drilling cost", "10 / 4 / 7", "BWP per hole / per kg / per m"],
            ["Distance R", "500", "m"],
            ["PPV K / β / limit", "1000 / 1.6 / 12.5", "site constants; limit in mm/s"],
            ["Airblast K_air / B_air / limit", "170 / 20 / 134", "site constants; limit in dB"],
            ["Rock factor A / RR n", "22 / 1.8", "fragmentation constants"],
            ["Target X50 / oversize threshold / allowed oversize", "120 / 500 / 10", "mm / mm / %"],
            ["Burden bounds", "2.5–4.5", "m"],
            ["Spacing:B, stemming:B, subdrill:B, stiffness bounds", "Interface defaults", "dimensionless engineering constraints"],
        ],
    )
    document.add_heading("8.2 Actions and outputs", level=2)
    add_steps(
        document,
        [
            "Enter geometry, explosive properties, costs, site constants, limits and engineering constraints.",
            "Select objective mode: Min Cost; Min Cost + Frag; or Min Cost + Frag + PPV/Air.",
            "Confirm weights and objective toggles.",
            "Select Compute KPIs to evaluate the current design.",
            "Review total cost in BWP, PPV in mm/s, airblast in dB, X50 in mm, oversize %, PF in kg/m³ and every constraint status.",
            "For a single constrained solution, use SLSQP and select Optimise.",
            "For trade-off exploration, choose Pareto and select Optimise; inspect multiple frontier rows.",
            "Export all/frontier rows as CSV or export the cost report.",
        ],
    )
    add_callout(
        document,
        "COST ASSUMPTION",
        "The interface labels cost in Botswana pula (BWP). Confirm that all entered rates use the same currency basis, tax basis and effective date before comparing scenarios.",
        "warning",
    )
    add_screenshot(
        document,
        "Figure 7 — Cost Optimisation input groups and controls.",
        ("cost", "optimisation"),
        "Insert Cost Optimisation after Compute KPIs, showing representative inputs, output KPI cards and constraint checks.",
    )
    add_screenshot(
        document,
        "Figure 7A — Cost Optimisation settings and engineering report.",
        ("cost", "optimisation", "results"),
        "Insert the Cost Optimisation engineering report after Compute KPIs using approved example values.",
    )

    # 9 GeoMotion
    document.add_heading("9. GeoMotion 3D", level=1)
    add_module_intro(
        document,
        "Move a charged-hole tie-up and optional one-metre mining block model through a reduced-order firing-event simulation, then study post-blast movement, ore loss, dilution, recovery and mixing.",
        "Evaluating the complete blast-movement workflow, timing sensitivity and ore-control concepts before a calibrated production model is available.",
        "Interactive 3D voxel model, movement/event diagnostics, ore-control KPIs, mixing matrix, movement CSV, full-resolution CSV.gz and result JSON.",
    )
    add_callout(
        document,
        "SYNTHETIC DEMONSTRATION / UNCALIBRATED — PLANNING ONLY",
        "Current movement, geology, grade, recovery, dilution and uncertainty outputs are synthetic unless their provenance explicitly states measured. Do not use GeoMotion for field execution, firing, dig-limit control, resource reporting or production decisions. It is not a detonation hydrocode or a validated site movement predictor.",
        "danger",
    )
    document.add_heading("9.1 Four-stage workflow", level=2)
    add_steps(
        document,
        [
            "Upload a delay-bearing charged-hole tie-up, or select Load 182-hole diamond demonstration.",
            "Optionally upload a measured 1 m mining block model. If none is supplied, the module creates a clearly labelled simulated block model.",
            "Review movement, explosive, rock, joint, loader and timing assumptions, then select Event physics baseline or Event physics + uncertainty realization.",
            "Select Run GeoMotion 3D, review the model provenance and validation messages, inspect ore-control outcomes, and export planning-only results.",
        ],
    )
    add_screenshot(
        document,
        "Figure 8 — GeoMotion 3D upload, assumptions and run workflow.",
        ("geomotion", "workflow"),
        "Insert an approved screenshot showing the four GeoMotion stages, tie-up input, optional block-model input and Run GeoMotion 3D control.",
    )

    document.add_heading("9.2 Required charged-hole tie-up", level=2)
    add_table(
        document,
        ["Column", "Required", "Unit / behaviour"],
        [
            ["Hole ID", "Recommended", "Unique text identifier; missing IDs can be generated"],
            ["X", "Yes", "Projected collar easting in m"],
            ["Y", "Yes", "Projected collar northing in m"],
            ["Z", "Yes", "Collar elevation in m RL"],
            ["Depth", "Yes", "Total drilled depth in m; must be at least 1 m"],
            ["Charge", "Yes", "Positive explosive mass in kg"],
            ["Delay", "Yes", "Unique cumulative nominal firing time in ms"],
        ],
    )
    document.add_paragraph("Example:")
    example = document.add_paragraph()
    example.style = document.styles["No Spacing"]
    run = example.add_run(
        "Hole ID,Depth,Charge,X,Y,Z,Delay\n"
        "N1,14.454,616.027,-5363.884,4603.971,678.454,8000"
    )
    run.font.name = "Aptos Mono"
    run.font.size = Pt(8.5)
    add_callout(
        document,
        "TIMING INTERPRETATION",
        "Delay values must be cumulative firing times, not inter-hole increments. GeoMotion preserves the original values for audit and simulates Delay minus the minimum Delay. Duplicate or missing delays are rejected; the module does not invent a timing pattern.",
        "note",
    )
    document.add_heading("9.3 Input validation", level=2)
    add_bullets(
        document,
        [
            "Blank or invalid X/Y/Z coordinates.",
            "Missing, non-positive or implausible depth and charge values.",
            "Missing or duplicate cumulative firing times.",
            "Duplicate Hole IDs and near-overlapping collars.",
            "Trailing empty rows, inferred floor RL and median nearest-hole spacing.",
            "At least three valid holes and a coordinate footprint spanning both X and Y are required.",
        ],
    )

    document.add_heading("9.4 Optional mining block model and measured datasets", level=2)
    add_table(
        document,
        ["Dataset", "Required fields / purpose"],
        [
            ["Mining block model", "CSV: X, Y, Z, Density; Block ID, Grade and Facies recommended. Optional Size X/Y/Z must each equal 1 m."],
            ["Geological structures/joints", "Registered measured structure data for future validated backend use."],
            ["Pre-blast / post-blast surface", "CSV containing X, Y, Z surface points."],
            ["Movement monitors", "CSV containing X, Y, Z and measured dX, dY, dZ."],
            ["Dig limits", "Polygon ID, vertex sequence, X, Y and destination."],
            ["Loader/MMU geometry", "Operational selectivity information for loader-scale review."],
        ],
    )
    document.add_paragraph(
        "A validated uploaded mining block model directly replaces simulated geology. Other optional files are registered and schema-validated, but registration alone does not mean their contents changed the simulation."
    )

    document.add_heading("9.5 Movement assumptions", level=2)
    add_table(
        document,
        ["Input", "Default", "Unit / purpose"],
        [
            ["Burden", "6", "m"],
            ["Spacing", "7", "m"],
            ["Hole diameter", "250", "mm"],
            ["Average stemming", "5.02", "m"],
            ["Subdrill", "1", "m"],
            ["Rock density", "2.35", "t/m³"],
            ["Powder factor", "0.92", "kg/m³"],
            ["Swell factor", "1.25", "ratio"],
            ["Cutoff grade", "12", "cpht"],
            ["Free-face azimuth", "180", "degrees"],
            ["Model cell", "1", "m; authoritative calculation resolution"],
            ["Relative energy", "1", "ratio"],
        ],
    )
    document.add_heading("9.6 Explosive, rock, joint and loader defaults", level=2)
    add_table(
        document,
        ["Input", "Default", "Unit / purpose"],
        [
            ["S135B density", "1,250.51", "kg/m³"],
            ["S135B RWS", "115", "%"],
            ["Nominal VOD / uncertainty", "4,500 / 1,000", "m/s"],
            ["Electronic timing scatter", "Enabled", "σ = 0.094 + 0.000345 × normalized delay, ms"],
            ["UCS / tensile strength", "120 / 10", "MPa"],
            ["Young's modulus", "55", "GPa"],
            ["Poisson ratio / damping", "0.24 / 0.28", "ratios"],
            ["Fragmentation index", "0.55", "0–1"],
            ["Joint dip / direction", "70 / 90", "degrees"],
            ["Joint spacing / persistence", "2.5 / 0.60", "m / 0–1"],
            ["Loader bucket", "100", "t"],
            ["Minimum mining unit", "5", "m"],
        ],
    )
    add_callout(
        document,
        "ASSUMPTION CONTROL",
        "Defaults are synthetic/site assumptions until replaced by signed-off mine files, laboratory measurements or manufacturer records. Record every changed value, source and approver with the exported result.",
        "warning",
    )

    document.add_heading("9.7 Engine modes and fallback", level=2)
    add_table(
        document,
        ["Mode/state", "Meaning"],
        [
            ["Event physics baseline", "Reduced-order timed event physics without a site-calibrated residual."],
            ["Event physics + uncertainty realization", "Event physics with repeatable sampled timing and VOD uncertainty; not a validated AI prediction."],
            ["Authoritative backend", "Computes contiguous 1 m voxels and can provide full-resolution gzip CSV export."],
            ["Coarse browser preview", "Used only if the GeoMotion backend returns 404/405; runs a clearly labelled 3 m preview and is not the full engine."],
        ],
    )

    document.add_heading("9.8 3D review controls", level=2)
    add_table(
        document,
        ["Control", "Options / purpose"],
        [
            ["Model view", "In-situ model; Movement timeline; Post-blast model"],
            ["Colour", "Ore/waste; Kimberlite facies; grade cpht; displacement; uncertainty; burden velocity; peak impulse"],
            ["Vectors", "Show sampled source-to-destination vectors"],
            ["Camera", "Perspective; Plan; Section; drag to orbit and scroll to zoom"],
            ["Z exaggeration", "1×, 2× or 3×"],
            ["Voxel seam", "Joined or 1%, 2%, 3% separation"],
            ["Section clip", "0–100% internal clipping"],
            ["Timeline", "Scrub from in-situ to post-blast and review event number, hole and actual firing time"],
        ],
    )
    add_screenshot(
        document,
        "Figure 8A — GeoMotion 3D movement model, legend and ore-control KPIs.",
        ("geomotion", "3d", "results"),
        "Insert an approved post-run screenshot showing the solid voxel model, selected colour legend, camera/timeline controls and KPI cards.",
    )

    document.add_heading("9.9 Outputs and interpretation", level=2)
    add_table(
        document,
        ["Output group", "Contents / units"],
        [
            ["Ore control", "Ore recovery %, ore loss %, dilution %, carat recovery %, predicted feed grade cpht and tonnes."],
            ["Movement", "Mean/P95 displacement m, mean heave m, maximum throw m, burden velocity m/s and peak impulse m/s."],
            ["Uncertainty", "Mean and P95 uncertainty in m, method and out-of-domain status."],
            ["Mixing matrix", "Ore→ore, ore→waste, waste→ore and waste→waste tonnes and percent of total."],
            ["Loader scale", "Recovery and dilution using the configured minimum mining unit."],
            ["Conservation", "Modelled tonnes, contained carats, remap collisions and mass/carat balance flags."],
            ["Events", "Nominal/actual time, timing error, charge, VOD, pressure proxy, energy, stemming effectiveness, burden velocity and released voxels."],
        ],
    )
    document.add_paragraph(
        "A zero mass-balance error confirms numerical conservation only; it does not prove that the predicted movement is accurate. Recovery and dilution become decision-relevant only after measured geology, dig limits and movement calibration are validated."
    )
    document.add_heading("9.10 Export formats", level=2)
    add_table(
        document,
        ["Action", "File / contents"],
        [
            ["Movement CSV", "geomotion_3d_movement_vectors_synthetic.csv; displayed blocks with source/destination, vector, velocity, geology, grade, tonnes, carats, provenance and planning notice."],
            ["Full 1 m CSV.gz", "geomotion_1m_full_resolution.csv.gz from the authoritative backend."],
            ["Result JSON", "geomotion_3d_synthetic_result.json; project notice, assumptions, validation, metrics, blocks, events, transport, provenance and remap."],
        ],
    )

    # 10 slope
    document.add_heading("10. Slope Stability", level=1)
    add_module_intro(
        document,
        "Screen a slope condition as likely Stable or Failure using a model or local factor-of-safety-style fallback.",
        "A rapid preliminary geotechnical screening that will be escalated for engineering assessment.",
        "Stable/Failure classification, stable probability, train/test accuracy, class balance and slope sketch.",
    )
    document.add_heading("10.1 Inputs", level=2)
    add_table(
        document,
        ["Input", "Range in interface", "Unit / meaning"],
        [
            ["H", "1–50", "m; slope height"],
            ["β", "5–80", "degrees; slope angle"],
            ["c", "1–200", "kPa; cohesion"],
            ["φ", "5–60", "degrees; friction angle"],
            ["γ", "14–28", "kN/m³; unit weight"],
            ["ru", "0–1", "dimensionless pore-pressure ratio"],
            ["B", "0–30", "m; sketch width only, not a model input"],
        ],
    )
    document.add_heading("10.2 File format and procedure", level=2)
    document.add_paragraph(
        "The upload accepts CSV, XLSX or XLS. Required model fields correspond to γ, c, φ, β, H, ru and Status. Status should identify Stable or Failure. Use clear ASCII headers where possible, for example gamma_kN_m3,c_kPa,phi_deg,beta_deg,H_m,ru,status."
    )
    add_steps(
        document,
        [
            "Load an approved labelled slope dataset, or use the configured default.",
            "Select Load & Predict to prepare the model and initialise parameters.",
            "Adjust H, β, c, φ, γ and ru for the case. B changes only the sketch.",
            "Review Stable/Failure and probability, plus train/test accuracy and class balance.",
            "Escalate uncertain, out-of-range or failure-prone cases for a full geotechnical assessment.",
        ],
    )
    add_callout(
        document,
        "FALLBACK BEHAVIOUR",
        "If the backend cannot provide a model response, the interface can silently display a local factor-of-safety-style estimate. Confirm model provenance before treating the probability as an ML result.",
        "warning",
    )
    add_screenshot(
        document,
        "Figure 9 — Slope Stability inputs and slope sketch.",
        ("slope", "stability", "results"),
        "Insert a Slope Stability sample result showing all parameter units, classification probability and slope sketch.",
    )

    # 11 backbreak
    document.add_heading("11. Back Break", level=1)
    add_module_intro(
        document,
        "Train a Random Forest on historical wall-control data and estimate back-break distance for an input scenario.",
        "Comparing perimeter-control scenarios and identifying variables associated with wall damage.",
        "Predicted back break, train/test R², feature importance and response surface.",
    )
    document.add_heading("11.1 CSV format", level=2)
    add_table(
        document,
        ["Typical column", "Unit / role"],
        [
            ["Burden", "m; model feature"],
            ["Spacing", "m; model feature"],
            ["Stemming", "m; model feature"],
            ["Powder Factor", "kg/m³; model feature"],
            ["Stiffness ratio", "dimensionless; model feature"],
            ["Backbreak", "m; required/automatically inferred target"],
        ],
    )
    document.add_heading("11.2 Procedure", level=2)
    add_steps(
        document,
        [
            "Select a representative historical CSV with numeric features and a back-break target.",
            "Select Predict Now.",
            "Review predicted back break in metres and train/test R².",
            "Adjust feature sliders within observed ranges; use Reset to Medians when needed.",
            "Review feature importance and choose two different surface axes.",
            "Use the result with Slope Stability and Flyrock screening; do not extrapolate beyond credible site ranges.",
        ],
    )
    add_screenshot(
        document,
        "Figure 10 — Back Break data and prediction controls.",
        ("back", "break", "results"),
        "Insert Back Break after running the approved sample dataset, with predicted value, R² and the response surface visible.",
    )

    # 12 flyrock
    document.add_heading("12. Flyrock (ML + Empirical)", level=1)
    add_module_intro(
        document,
        "Estimate flyrock throw distance with a trained model and compare it with an automatically selected empirical method.",
        "Screening throw risk and informing a separately approved exclusion-zone assessment.",
        "Predicted distance, empirical distance, train/test R², feature importance and response surface.",
    )
    document.add_heading("12.1 CSV format", level=2)
    add_table(
        document,
        ["Typical column", "Unit / role"],
        [
            ["Hole diameter", "mm; feature"],
            ["Burden / Spacing / Stemming / Bench height", "m; features"],
            ["Charge per delay", "kg; feature"],
            ["Powder factor", "kg/m³; feature"],
            ["Rock density", "t/m³; feature"],
            ["SDoB", "dimensionless scaled depth of burial; optional empirical feature"],
            ["Flyrock distance / Throw", "m; target"],
        ],
    )
    document.add_heading("12.2 Procedure", level=2)
    add_steps(
        document,
        [
            "Load a representative CSV containing a recognised flyrock/throw target.",
            "Select Predict.",
            "Review the ML prediction and the empirical estimate in metres, including the selected empirical method.",
            "Check train/test R² and feature importance.",
            "Adjust input sliders only within supported data ranges.",
            "Select two different surface axes or Redraw surface to inspect sensitivity.",
            "Compare the output with the approved site limit and exclusion-zone method; apply the more conservative qualified assessment.",
        ],
    )
    add_callout(
        document,
        "EXCLUSION-ZONE WARNING",
        "The flyrock estimate is not an automatic clearance distance. Exclusion zones must be established by qualified personnel using legislation, mine procedures, blast-specific hazards, empirical evidence and conservative controls.",
        "danger",
    )
    add_screenshot(
        document,
        "Figure 11 — Flyrock ML and empirical analysis controls.",
        ("flyrock", "results"),
        "Insert a successful Flyrock run showing predicted and empirical distance, model quality and response surface.",
    )

    # 13 formats
    document.add_heading("13. File formats and output handling", level=1)
    add_table(
        document,
        ["Module", "Accepted input", "Exports / visible outputs"],
        [
            ["Data Manager", "CSV", "CSV; nominal XLSX; site-model JSON; plots"],
            ["Prediction", "Active dataset; uploaded CSV through Data Manager", "prediction_result.csv; printable HTML report"],
            ["Feature Importance", "Active dataset; uploaded CSV through Data Manager", "On-screen diagnostics and plots"],
            ["Parameter Optimisation", "Active dataset; uploaded CSV through Data Manager", "param_surface.csv; printable HTML report"],
            ["Cost Optimisation", "Manual numeric inputs", "Cost report HTML; Pareto all/frontier CSV"],
            ["GeoMotion 3D", "Charged-hole CSV; optional 1 m block-model CSV and measured datasets", "Movement CSV; full-resolution CSV.gz; result JSON; 3D model and KPIs"],
            ["Slope Stability", "CSV, XLSX, XLS", "On-screen classification and sketch"],
            ["Back Break", "CSV", "On-screen prediction, importance and surface"],
            ["Flyrock", "CSV", "On-screen ML/empirical prediction, importance and surface"],
        ],
    )
    document.add_heading("13.1 CSV rules", level=2)
    add_bullets(
        document,
        [
            "Use one header row and one observation or hole per subsequent row.",
            "Use a period as the decimal separator and do not include thousands separators in numeric cells.",
            "Put units in headers, not in numeric values.",
            "Preserve exact column names expected by the module or verify automatic mapping.",
            "Remove merged cells, formulas returning errors, embedded notes and hidden totals before import.",
            "Retain the original source file and create a working copy for cleaning.",
        ],
    )
    document.add_heading("13.2 Printable HTML reports", level=2)
    document.add_paragraph(
        "Prediction, Parameter Optimisation and Cost Optimisation generate printable HTML. The report opens in a browser tab; use the browser Print command to print or Save as PDF. GeoMotion exports structured CSV/CSV.gz and JSON rather than a printable HTML report."
    )
    document.add_heading("13.3 Record retention", level=2)
    add_bullets(
        document,
        [
            "Record source dataset filename, revision/date, selected combined dataset and any filters.",
            "Record all site constants, limits, optimisation weights and manual overrides.",
            "Retain exported results with qualified review, approval and operational records under the site's document-control policy.",
            "Do not include authentication tokens, one-time codes or unnecessary personal data in exports.",
        ],
    )

    # 14 troubleshooting
    document.add_heading("14. Troubleshooting", level=1)
    add_table(
        document,
        ["Symptom", "Likely cause", "Action"],
        [
            ["Email not authorised", "Address is not in frontend/backend allowlists or deployed lists differ.", "Confirm spelling and ask the access owner to verify both deployed allowlists."],
            ["No verification code", "Mail delay, spam filtering or InstantDB configuration.", "Wait briefly, check junk/quarantine, retry once and contact support without sharing codes."],
            ["Missing InstantDB app ID", "Authentication environment variable is absent.", "Administrator must configure VITE_INSTANTDB_APP_ID."],
            ["Backend offline/error", "API URL, service, CORS, authentication or network issue.", "Check status indicator and approved service status; provide timestamp and module to support."],
            ["CSV loads incorrectly", "Wrong delimiter, quoted fields, blank headers or unit text.", "Save as UTF-8 comma-separated CSV with one header row."],
            ["Prediction inputs remain loading", "Metadata endpoint or dataset is unavailable.", "Confirm backend and combined dataset; refresh after service recovery."],
            ["No ML output", "Model assets unavailable or dataset cannot train a model.", "Use empirical output cautiously and ask the model owner to check assets/data."],
            ["Poor/negative test R²", "Weak, small or non-representative dataset.", "Do not rely on the model; improve data and obtain technical review."],
            ["Parameter optimisation times out", "Search workload or backend response time.", "Retry once, use valid distinct axes and review dataset size/ranges."],
            ["Constraint shows Check", "Current or optimised design violates a configured limit.", "Review the specific ratio/limit; do not approve until resolved by qualified personnel."],
            ["GeoMotion rejects the tie-up", "Required X/Y/Z/depth/charge/delay data are missing, invalid, or delays are duplicated.", "Correct the source charged-hole file; use unique cumulative Delay values for every valid hole."],
            ["GeoMotion shows coarse 3 m preview", "The upgraded backend returned 404/405.", "Treat it only as a browser preview; deploy/restore the authoritative backend before expecting 1 m physics or full-resolution export."],
            ["GeoMotion block model rejected", "Required fields are missing or block dimensions are not 1 m.", "Provide X,Y,Z,Density and resample Size X/Y/Z to 1 m; include Block ID, Grade and Facies where available."],
            ["Printable report does not open", "Browser blocked the new tab.", "Allow pop-ups for the approved site or use the downloaded HTML fallback."],
            ["Excel export will not open correctly", "Current Data Manager export is CSV content with .xlsx name.", "Use Export Filtered → CSV; convert to XLSX in spreadsheet software if required."],
        ],
    )
    document.add_heading("14.1 Information to provide to support", level=2)
    add_bullets(
        document,
        [
            "Date/time, module, browser and whether the backend status appeared online.",
            "Dataset filename and column headers only, unless the support channel is approved for operational data.",
            "The exact visible error message and steps immediately before it.",
            "A redacted screenshot with secrets, codes and sensitive production values removed.",
            "Never send passwords, one-time codes, refresh tokens or API keys.",
        ],
    )
    add_callout(
        document,
        "OWNER ACTION REQUIRED — SUPPORT DETAILS",
        "Insert the approved service desk email/phone, escalation path, operating hours, incident severity rules and production status-page address.",
        "action",
    )

    # 15 glossary
    document.add_heading("15. Glossary", level=1)
    add_table(
        document,
        ["Term", "Definition"],
        [
            ["Airblast", "Pressure wave transmitted through air and represented here as a dB level."],
            ["Back break", "Rock breakage extending behind the intended final wall or last row."],
            ["Empirical model", "Equation fitted or calibrated from observed practice rather than a learned black-box model."],
            ["Flyrock", "Rock projected beyond the intended blast area."],
            ["cpht", "Carats per hundred tonnes; grade unit used in GeoMotion."],
            ["Dilution", "Waste entering material classified and routed as ore."],
            ["Facies", "Geological material subdivision with distinct properties."],
            ["Kuz–Ram", "Empirical fragmentation model used to estimate a characteristic mean size."],
            ["Machine learning (ML)", "Model trained from historical input/output observations."],
            ["Minimum mining unit (MMU)", "Operational selectivity scale used to calculate loader-scale recovery and dilution."],
            ["Ore loss", "Source ore ending in a waste-classified destination."],
            ["Pareto frontier", "Solutions for which one objective cannot improve without worsening another."],
            ["PCA", "Principal component analysis; transforms correlated variables into components explaining variance."],
            ["Permutation importance", "Model performance loss when one feature is shuffled."],
            ["Rosin–Rammler (RR)", "Particle-size distribution model using characteristic size and uniformity n."],
            ["Scaled burden", "Burden normalised by a charge-related term for screening confinement/throw risk."],
            ["Scaled distance", "Distance normalised by charge per delay for vibration/airblast relationships."],
            ["Surrogate model", "Faster approximation used to explore or optimise an expensive response."],
            ["Voxel", "Three-dimensional material cell; GeoMotion's authoritative source model uses contiguous 1 m cells."],
        ],
    )

    # 16 owner completion
    document.add_heading("16. Owner completion checklist", level=1)
    document.add_paragraph(
        "Complete and approve every item below before issuing this manual as a controlled production document."
    )
    add_table(
        document,
        ["Item", "Required owner action", "Complete"],
        [
            ["Production identity", f"Add organisation logo/name, software build/release and document identifier; verify {DEPLOYMENT_URL}.", "☐"],
            ["Approval", "Add document owner, technical reviewer, qualified blasting reviewer and release signatures.", "☐"],
            ["Access list", "Verify deployed frontend/backend allowlists and define access review/revocation process.", "☐"],
            ["Support", "Add service desk and escalation details.", "☐"],
            ["Site constants", "Document approved PPV, airblast, fragmentation and flyrock limits and calibrated constants.", "☐"],
            ["Coordinate system", "State the approved X/Y/Z coordinate reference system and units.", "☐"],
            ["Cost basis", "Confirm currency, effective date, tax basis and cost-rate ownership.", "☐"],
            ["Model governance", "Add model/dataset owners, approved datasets, GeoMotion calibration status, validation thresholds and retraining/change controls.", "☐"],
            ["GeoMotion data", "Define approved tie-up, block-model, survey, movement-monitor, dig-limit and reconciliation sources with sign-off responsibilities.", "☐"],
            ["Figures", "Insert approved deployment screenshots in each labelled figure placeholder.", "☐"],
            ["Excel limitation", "Correct or formally document the nominal XLSX export behaviour.", "☐"],
            ["Training", "Add required competence, induction and refresher requirements.", "☐"],
            ["Retention", "Add the applicable document and blast-record retention period.", "☐"],
        ],
    )
    add_callout(
        document,
        "FINAL RELEASE GATE",
        "A qualified blasting reviewer must confirm that the manual's limits, terminology, units, figures and workflows match the deployed software and local procedures. GeoMotion must remain planning-only until mine data, calibration and governance acceptance criteria are approved.",
        "danger",
    )

    return document


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build_manual()
    document.core_properties.title = "AI Blasting Suite User Manual"
    document.core_properties.subject = "Operating instructions, module inputs, outputs, formats and units"
    document.core_properties.author = "AI Blasting Suite documentation"
    document.core_properties.keywords = "blasting, user manual, prediction, optimisation, GeoMotion 3D, ore movement"
    document.core_properties.comments = "Draft for owner and qualified blasting review."
    document.save(OUTPUT)
    print(f"Generated {OUTPUT}")


if __name__ == "__main__":
    main()
